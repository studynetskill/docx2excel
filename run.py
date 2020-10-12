import docx
from docx import Document  # 导入读取docx库
import re  # 正则匹配ABC
from openpyxl import Workbook  # 写xlsx库
import os  # 读取文件


from ReadDocx import ReadDocx
from setting import TEMPLATE_FILE, DOCX_PATH


# file_path = r'F:\Python\doc2xls\code\docx'
# doc_path = 'F:\\Python\\doc2xls\\code\\申请表汇总\\'
# path = '123.docx'  # 文件路径
# test_path = '测试数据.docx'  # 测试数据
# document = Document(path)  # 读入文件
# tables = document.tables  # 获取文件中的表格集


# 读取docx文件
def getDocxFile(docx_path):
    '''
    :param docx_path: docx文件所在目录
    :return docx_list: 返回docx文件名数组
    '''
    docx_list = []
    for i in os.listdir(docx_path):
        # 找出文件中以.docx结尾并且不以~$开头的文件（~$是为了排除临时文件的）
        if i.endswith('.docx') and not i.startswith('~$'):
            # print(i)
            docx_list.append(os.path.join(docx_path, i))
    return docx_list


# 根据模板里的大写字母A-Z匹配数据所在位置，并返回位置的数组
def template_pattern(table_data):
    '''
    :param table_data: 模板数组，用A，B，AA等标记要匹配的位置
    :return model_pattern: 返回一个包含匹配后标记所在位子的数组
    '''
    model_pattern = []
    pattern = re.compile(r"^[A-Z]{1,2}$")
    for i, cell in enumerate(table_data):
        # print(i, cell)
        res = pattern.match(cell)
        # print(res)
        if res is not None:
            model_pattern.append(i)
    # print(model_pattern)
    if len(model_pattern):
        return model_pattern
    return None


# 根据匹配的位置获取指定数据
def math_docx(table_data, template_pattern):
    '''
    :param table_data: 要匹配的原始数据数组
        template_pattern：根据这个数组里所在的标记位置来匹配
    '''
    math_data = []
    for i in template_pattern:
        # print(i)
        math_data.append(table_data[i])
    return math_data


# 将数据写入xlsx
def write_xlxs(data, result_xlsx="docx2xlsx收集结果.xlsx"):
    '''
    :param data:数组列表，每个元素数组为一行
        result_xlsx:默认结果储存在这个xlsx
    :return: None
    '''
    wb = Workbook()
    # 激活默认表单
    ws = wb.active
    # 以行为单位输入
    for row in data:
        ws.append(row)
    # 保存文件
    wb.save(result_xlsx)


def main(template_file, docx_path):
    '''
    主程序入口
    :param template_file: 模板文件
        docx_path: docx文件路径
    '''

    # 模板标记位置获取
    template_doc = ReadDocx(template_file)
    temp_data = template_doc.read_table()
    data_pattern = template_pattern(temp_data)
    # print(data_pattern)
    if data_pattern is None:
        print('请检查模板文件是否正确标记')
        return
    docx_list = getDocxFile(docx_path)
    # print(docx_list)
    excel_data = []
    # 依次读取docx并将每一个文件为数组的一个元素
    for docx_file in docx_list:
        docx_data = ReadDocx(docx_file)
        table_data = docx_data.read_table()
        mached_data = math_docx(table_data, data_pattern)
        # print(mached_data)
        excel_data.append(mached_data)
    # print(excel_data)
    # 写入xlsx
    write_xlxs(excel_data)


if __name__ == "__main__":
    print('开始收集...请耐心等候')
    main(TEMPLATE_FILE, DOCX_PATH)
    print('收集完成!')
